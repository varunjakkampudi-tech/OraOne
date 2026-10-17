"""Contact import for leads: parse, validate and de-duplicate CSV / XLSX uploads.

Two steps, both org-scoped by the caller:

* ``prepare_import`` + ``classify_rows`` -- read and validate the file and
  mark each row valid / duplicate / invalid. Reads existing leads only.
  Writes nothing.
* ``save_valid_rows`` -- saves only the rows classified as valid, tagged with
  source ``import``, the upload id and the uploader.

Files are checked before any parser runs: size, extension, magic bytes, the
zip structure of XLSX files and macro parts. Contact details are never logged.
"""

from __future__ import annotations

import csv
import io
import os
import re
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from typing import Any, Optional

from fastapi import UploadFile
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database.models.lead import Lead, LeadStatus
from app.services import workflow_triggers
from app.services.lead_service import score_lead

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_DATA_ROWS = 5000
MAX_SCANNED_ROWS = 20000
MAX_UNPACKED_XLSX_BYTES = 50 * 1024 * 1024

ROW_VALID = "valid"
ROW_DUPLICATE_EXISTING = "duplicate_existing"
ROW_DUPLICATE_IN_FILE = "duplicate_in_file"
ROW_INVALID = "invalid"

_ZIP_MAGIC = b"PK\x03\x04"
_EXTENSION_FORMATS = {".csv": "csv", ".xlsx": "xlsx"}
_MACRO_EXTENSIONS = {".xlsm", ".xltm", ".xlam", ".xlsb"}

# Header spellings (normalised: lowercase, letters and digits only) -> field.
_HEADER_ALIASES: dict[str, str] = {}
for _field, _aliases in {
    "name": ("name", "fullname", "contactname", "leadname"),
    "phone": (
        "phone",
        "phonenumber",
        "mobile",
        "mobilenumber",
        "tel",
        "telephone",
        "cell",
        "contactnumber",
    ),
    "email": ("email", "emailaddress", "mail", "emailid"),
    "company": (
        "company",
        "companyname",
        "organisation",
        "organization",
        "business",
        "businessname",
    ),
    "notes": ("notes", "note", "comments", "comment"),
    "consent": ("consent", "marketingconsent", "optin", "consentstatus"),
}.items():
    for _alias in _aliases:
        _HEADER_ALIASES[_alias] = _field

_CONSENT_YES = {"yes", "y", "true", "1", "granted", "opted in", "opt in", "consented"}
_CONSENT_NO = {"no", "n", "false", "0", "denied", "opted out", "opt out", "declined"}

#: Recorded on leads whose consent came from a CSV cell. The CSV is self-attested:
#: the importer asserts consent, the system did not capture it. Nothing is invented:
#: a blank cell stays "unknown" and gets no source.
IMPORT_CONSENT_SOURCE = "import_csv"
IMPORTED_CONSENT_NOTE = "Imported consent, confirm the source"

_PHONE_CHARS_RE = re.compile(r"^\+?[0-9 ().\-]+$")
_EMAIL_RE = re.compile(
    r"^[a-z0-9._%+\-']+@[a-z0-9](?:[a-z0-9\-]{0,61}[a-z0-9])?(?:\.[a-z0-9](?:[a-z0-9\-]{0,61}[a-z0-9])?)+$"
)


class LeadImportError(Exception):
    """A file-level problem. ``message`` is plain language and safe to show."""

    def __init__(self, status_code: int, message: str) -> None:
        super().__init__(message)
        self.status_code = status_code
        self.message = message


@dataclass
class ParsedRow:
    row_number: int
    status: str = ROW_VALID
    reason: Optional[str] = None
    name: Optional[str] = None
    phone: Optional[str] = None  # stored form, e.g. "+14155550123"
    phone_key: Optional[str] = None  # digits only, used for matching
    email: Optional[str] = None
    company: Optional[str] = None
    notes: Optional[str] = None
    consent: Optional[str] = None


@dataclass
class ParsedImport:
    file_name: str
    file_format: str
    columns: list[tuple[str, Optional[str]]]
    rows: list[ParsedRow]
    warnings: list[str] = field(default_factory=list)


# ───────────────────────── file-level checks ─────────────────────────


async def read_upload(file: UploadFile) -> bytes:
    """Read at most MAX_FILE_BYTES + 1 bytes so an oversized upload is never buffered whole."""
    data = await file.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise LeadImportError(
            413,
            "This file is larger than 5 MB. Split it into smaller files and try again.",
        )
    if not data:
        raise LeadImportError(422, "This file is empty.")
    return data


def safe_file_name(filename: Optional[str]) -> str:
    base = os.path.basename((filename or "").replace("\\", "/")).strip()
    return base[:120] or "contacts"


def detect_format(filename: str, data: bytes) -> str:
    """Decide the format from the extension AND the content. Neither alone is trusted."""
    suffix = os.path.splitext((filename or "").lower())[1]
    if suffix in _MACRO_EXTENSIONS:
        raise LeadImportError(
            415,
            "Macro-enabled or binary spreadsheets aren't accepted. Save the file as .xlsx or .csv and try again.",
        )
    fmt = _EXTENSION_FORMATS.get(suffix)
    if fmt is None:
        raise LeadImportError(415, "Only .csv and .xlsx files can be imported.")

    if fmt == "xlsx":
        if not data.startswith(_ZIP_MAGIC):
            raise LeadImportError(415, "This file isn't a valid .xlsx workbook.")
        _inspect_xlsx(data)
    else:
        if data.startswith(_ZIP_MAGIC) or b"\x00" in data[:4096]:
            raise LeadImportError(415, "This doesn't look like a plain-text CSV file.")
    return fmt


def _inspect_xlsx(data: bytes) -> None:
    """Reject macro parts, encryption and zip bombs before openpyxl opens the file."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise LeadImportError(415, "This file isn't a valid .xlsx workbook.") from exc

    with zf:
        infos = zf.infolist()
        names = {info.filename.lower() for info in infos}
        if any(name.endswith("vbaproject.bin") for name in names):
            raise LeadImportError(
                415,
                "Workbooks containing macros aren't accepted. Remove the macros and try again.",
            )
        if "xl/workbook.xml" not in names:
            raise LeadImportError(415, "This file isn't a valid .xlsx workbook.")
        if any(info.flag_bits & 0x1 for info in infos):
            raise LeadImportError(
                422,
                "This workbook is password protected. Remove the password and try again.",
            )
        if sum(info.file_size for info in infos) > MAX_UNPACKED_XLSX_BYTES:
            raise LeadImportError(422, "This workbook is too large once unpacked.")
        for info in infos:
            if info.filename.lower() == "[content_types].xml":
                content_types = zf.read(info).lower()
                if b"macroenabled" in content_types:
                    raise LeadImportError(
                        415,
                        "Workbooks containing macros aren't accepted. Remove the macros and try again.",
                    )


# ───────────────────────── parsing ─────────────────────────


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        # Excel stores whole numbers (phones, ids) as floats; keep them as digits.
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def parse_csv(data: bytes) -> list[tuple[int, list[str]]]:
    try:
        text = data.decode("utf-8-sig")  # strips a leading BOM if present
    except UnicodeDecodeError as exc:
        raise LeadImportError(
            422, "This CSV isn't UTF-8 text. Save it as 'CSV UTF-8' and try again."
        ) from exc

    first_line = text.split("\n", 1)[0]
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)

    rows: list[tuple[int, list[str]]] = []
    try:
        for number, cells in enumerate(reader, start=1):
            if number > MAX_SCANNED_ROWS:
                raise LeadImportError(422, _too_many_rows_message())
            if not any(c.strip() for c in cells):
                continue
            rows.append((number, [c.strip() for c in cells]))
            if len(rows) > MAX_DATA_ROWS + 1:  # +1 for the header row
                raise LeadImportError(422, _too_many_rows_message())
    except csv.Error as exc:
        raise LeadImportError(
            422,
            "This CSV couldn't be read. Check for unbalanced quote marks and try again.",
        ) from exc
    return rows


def parse_xlsx(data: bytes) -> list[tuple[int, list[str]]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise LeadImportError(
            422, "Couldn't open this workbook. Save it again as .xlsx and try again."
        ) from exc

    try:
        if not wb.worksheets:
            raise LeadImportError(422, "This workbook has no sheets.")
        sheet = wb.worksheets[0]
        rows: list[tuple[int, list[str]]] = []
        for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            if number > MAX_SCANNED_ROWS:
                raise LeadImportError(422, _too_many_rows_message())
            cells = [_cell_text(v) for v in values]
            if not any(cells):
                continue
            rows.append((number, cells))
            if len(rows) > MAX_DATA_ROWS + 1:
                raise LeadImportError(422, _too_many_rows_message())
        return rows
    except LeadImportError:
        raise
    except Exception as exc:
        raise LeadImportError(
            422,
            "Couldn't read the first sheet of this workbook. Save it again and try again.",
        ) from exc
    finally:
        wb.close()


def _too_many_rows_message() -> str:
    return (
        f"This file has more than {MAX_DATA_ROWS:,} contact rows. "
        "Split it into smaller files and try again."
    )


def _field_for_header(header: str) -> Optional[str]:
    return _HEADER_ALIASES.get(re.sub(r"[^a-z0-9]", "", header.lower()))


# ───────────────────────── field validation ─────────────────────────


def check_phone(raw: str) -> tuple[Optional[str], Optional[str], Optional[str]]:
    """Return (stored, match_key, error). Digits are matched; '+' is kept when given."""
    text = raw.strip()
    if not text:
        return None, None, None
    if not _PHONE_CHARS_RE.match(text):
        return (
            None,
            None,
            "Phone number contains letters or symbols that aren't allowed",
        )
    international = text.startswith("+") or text.startswith("00")
    digits = re.sub(r"\D", "", text)
    if text.startswith("00"):
        digits = digits[2:]
    if not 7 <= len(digits) <= 15:
        return None, None, "Phone number needs 7 to 15 digits"
    if international and digits.startswith("0"):
        return (
            None,
            None,
            "International phone number can't start with 0 after the country code",
        )
    stored = f"+{digits}" if international else digits
    return stored, digits, None


def check_email(raw: str) -> tuple[Optional[str], Optional[str]]:
    """Return (normalised_lowercase, error)."""
    text = raw.strip().lower()
    if not text:
        return None, None
    local = text.split("@", 1)[0]
    if (
        len(text) > 255
        or not _EMAIL_RE.match(text)
        or ".." in text
        or local.startswith(".")
        or local.endswith(".")
    ):
        return None, "Email address doesn't look valid"
    return text, None


def check_consent(raw: str) -> tuple[Optional[str], Optional[str]]:
    """Return (granted | denied | unknown, error). A blank cell means unknown."""
    text = raw.strip().lower()
    if not text:
        return "unknown", None
    if text in _CONSENT_YES:
        return "granted", None
    if text in _CONSENT_NO:
        return "denied", None
    return None, "Consent must be yes, no, or left blank"


def _validate_row(
    row_number: int, cells: list[str], index_by_field: dict[str, int]
) -> ParsedRow:
    def cell(name: str) -> str:
        idx = index_by_field.get(name)
        if idx is None or idx >= len(cells):
            return ""
        return cells[idx]

    problems: list[str] = []
    name = cell("name").strip()
    if len(name) > 160:
        problems.append("Name is longer than 160 characters")

    phone, phone_key, phone_error = check_phone(cell("phone"))
    if phone_error:
        problems.append(phone_error)

    email, email_error = check_email(cell("email"))
    if email_error:
        problems.append(email_error)

    company = cell("company").strip()
    if len(company) > 200:
        problems.append("Company name is longer than 200 characters")

    notes = cell("notes").strip()
    if len(notes) > 4000:
        problems.append("Notes are longer than 4,000 characters")

    consent, consent_error = check_consent(cell("consent"))
    if consent_error:
        problems.append(consent_error)

    if not phone and not email and not problems:
        problems.append(
            "Add a phone number or an email address so this contact can be reached"
        )

    row = ParsedRow(
        row_number=row_number,
        name=name or None,
        phone=phone,
        phone_key=phone_key,
        email=email,
        company=company or None,
        notes=notes or None,
        consent=consent,
    )
    if problems:
        row.status = ROW_INVALID
        row.reason = "; ".join(problems)
    return row


def prepare_import(data: bytes, filename: str) -> ParsedImport:
    """Parse and validate an upload. Pure: no database access. Raises LeadImportError."""
    fmt = detect_format(filename, data)
    table = parse_csv(data) if fmt == "csv" else parse_xlsx(data)
    if not table:
        raise LeadImportError(422, "This file has no rows.")

    _, header = table[0]
    data_rows = table[1:]
    if not data_rows:
        raise LeadImportError(422, "This file has column headings but no contact rows.")

    columns: list[tuple[str, Optional[str]]] = []
    index_by_field: dict[str, int] = {}
    for idx, heading in enumerate(header):
        field_name = _field_for_header(heading) if heading else None
        if field_name is not None and field_name in index_by_field:
            field_name = None  # first matching column wins
        if field_name is not None:
            index_by_field[field_name] = idx
        columns.append((heading, field_name))

    warnings: list[str] = []
    if "email" not in index_by_field and "phone" not in index_by_field:
        warnings.append(
            "No email or phone column was found, so every row will be skipped. "
            "Name a column 'Email' or 'Phone'."
        )

    rows = [_validate_row(number, cells, index_by_field) for number, cells in data_rows]
    return ParsedImport(
        file_name=safe_file_name(filename),
        file_format=fmt,
        columns=columns,
        rows=rows,
        warnings=warnings,
    )


# ───────────────────────── duplicate checks + save ─────────────────────────


async def _existing_keys(
    session: AsyncSession,
    *,
    organization_id: uuid.UUID,
    emails: set[str],
    phone_keys: set[str],
) -> tuple[set[str], set[str]]:
    """Return (matching emails, matching phone keys) already held by this organization."""
    if not emails and not phone_keys:
        return set(), set()

    clauses = []
    if emails:
        clauses.append(func.lower(Lead.email).in_(sorted(emails)))
    if phone_keys:
        clauses.append(
            func.regexp_replace(Lead.phone, r"\D", "", "g").in_(sorted(phone_keys))
        )

    result = await session.execute(
        select(Lead.email, Lead.phone)
        .where(Lead.organization_id == organization_id)
        .where(Lead.deleted_at.is_(None))
        .where(or_(*clauses))
    )
    found_emails: set[str] = set()
    found_phones: set[str] = set()
    for email, phone in result.all():
        if email and email.strip().lower() in emails:
            found_emails.add(email.strip().lower())
        key = re.sub(r"\D", "", phone or "")
        if key and key in phone_keys:
            found_phones.add(key)
    return found_emails, found_phones


async def classify_rows(
    session: AsyncSession, *, organization_id: uuid.UUID, parsed: ParsedImport
) -> None:
    """Mark valid rows as duplicates of existing leads or of earlier rows in the file."""
    candidates = [r for r in parsed.rows if r.status == ROW_VALID]
    emails = {r.email for r in candidates if r.email}
    phone_keys = {r.phone_key for r in candidates if r.phone_key}
    existing_emails, existing_phones = await _existing_keys(
        session, organization_id=organization_id, emails=emails, phone_keys=phone_keys
    )

    seen_emails: set[str] = set()
    seen_phones: set[str] = set()
    for row in candidates:
        if (row.email and row.email in existing_emails) or (
            row.phone_key and row.phone_key in existing_phones
        ):
            row.status = ROW_DUPLICATE_EXISTING
            row.reason = "Already a lead in your organization"
            continue
        if (row.email and row.email in seen_emails) or (
            row.phone_key and row.phone_key in seen_phones
        ):
            row.status = ROW_DUPLICATE_IN_FILE
            row.reason = "Repeats an earlier row in this file"
            continue
        if row.email:
            seen_emails.add(row.email)
        if row.phone_key:
            seen_phones.add(row.phone_key)


def summarise(parsed: ParsedImport) -> dict[str, int]:
    counts = {
        ROW_VALID: 0,
        ROW_DUPLICATE_EXISTING: 0,
        ROW_DUPLICATE_IN_FILE: 0,
        ROW_INVALID: 0,
    }
    for row in parsed.rows:
        counts[row.status] += 1
    return counts


async def save_valid_rows(
    session: AsyncSession,
    *,
    organization_id: uuid.UUID,
    project_id: uuid.UUID,
    user_id: uuid.UUID,
    parsed: ParsedImport,
) -> tuple[uuid.UUID, int]:
    """Add a Lead for each valid row. Caller commits. Returns (upload_id, saved_count)."""
    upload_id = uuid.uuid4()
    imported_at = datetime.now(timezone.utc)
    leads: list[Lead] = []
    for row in parsed.rows:
        if row.status != ROW_VALID:
            continue
        score, temperature = score_lead(
            email=row.email,
            phone=row.phone,
            company=row.company,
            intent=None,
            message=None,
        )
        extra: dict[str, Any] = {
            "consent": row.consent or "unknown",
            "import": {
                "upload_id": str(upload_id),
                "uploaded_by": str(user_id),
                "row": row.row_number,
            },
        }
        if row.consent in ("granted", "denied"):
            # The CSV is self-attested, so the record says where the value came from.
            extra["consent_source"] = IMPORT_CONSENT_SOURCE
            extra["consent_recorded_at"] = imported_at.isoformat()
        if row.consent == "granted":
            extra["consent_note"] = IMPORTED_CONSENT_NOTE
        if row.notes:
            extra["notes"] = row.notes
        leads.append(
            Lead(
                organization_id=organization_id,
                project_id=project_id,
                name=row.name,
                email=row.email,
                phone=row.phone,
                company=row.company,
                source="import",
                status=LeadStatus.new,
                score=score,
                temperature=temperature,
                extra=extra,
            )
        )
    session.add_all(leads)
    await session.flush()
    for lead in leads:
        await workflow_triggers.fire_lead_created(session, lead)
    return upload_id, len(leads)
